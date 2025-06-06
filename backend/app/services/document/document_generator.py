"""
Document Generation Service
Handles generation of HTML, PDF, and PowerPoint documents
"""
import logging
import io
import tempfile
import os
from typing import Dict, Any, Optional, Union
from pathlib import Path
from datetime import datetime
from dataclasses import dataclass
from enum import Enum

logger = logging.getLogger(__name__)


class DocumentFormat(Enum):
    """Supported document formats"""
    HTML = "html"
    PDF = "pdf"
    PPTX = "pptx"


class DocumentType(Enum):
    """Document types for templates"""
    RFP_ANALYSIS = "rfp_analysis"
    PROPOSAL = "proposal"
    RFP = "rfp"
    REPORT = "report"


@dataclass
class GeneratedDocument:
    """Generated document result"""
    content: bytes
    filename: str
    format: DocumentFormat
    content_type: str
    size: int
    generated_at: datetime


@dataclass
class DocumentTemplate:
    """Document template configuration"""
    name: str
    type: DocumentType
    format: DocumentFormat
    template_path: str
    description: str
    variables: Dict[str, str]


class DocumentGenerator:
    """Document generation service"""
    
    def __init__(self):
        self.template_dir = Path(__file__).parent / "templates"
        self.template_dir.mkdir(exist_ok=True)
        
        # Initialize template engines
        self._setup_template_engines()
        
        # Content type mappings
        self.content_types = {
            DocumentFormat.HTML: "text/html",
            DocumentFormat.PDF: "application/pdf",
            DocumentFormat.PPTX: "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        }
    
    def _setup_template_engines(self):
        """Initialize template engines"""
        try:
            # Jinja2 for HTML templates
            from jinja2 import Environment, FileSystemLoader
            self.jinja_env = Environment(
                loader=FileSystemLoader(str(self.template_dir)),
                autoescape=True
            )
            
            # Add custom filters
            self.jinja_env.filters['currency'] = self._currency_filter
            self.jinja_env.filters['datetime'] = self._datetime_filter
            self.jinja_env.filters['percentage'] = self._percentage_filter
            
            logger.info("Template engines initialized successfully")
            
        except ImportError as e:
            logger.error(f"Failed to initialize template engines: {e}")
            raise
    
    async def generate_document(
        self,
        template_name: str,
        data: Dict[str, Any],
        format: DocumentFormat,
        filename: Optional[str] = None
    ) -> GeneratedDocument:
        """
        Generate document from template and data
        
        Args:
            template_name: Name of the template to use
            data: Data to populate the template
            format: Output format (HTML, PDF, PPTX)
            filename: Optional custom filename
            
        Returns:
            Generated document with content and metadata
        """
        try:
            logger.info(f"Generating {format.value} document from template: {template_name}")
            
            # Generate filename if not provided
            if not filename:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"{template_name}_{timestamp}.{format.value}"
            
            # Generate based on format
            if format == DocumentFormat.HTML:
                content = await self._generate_html(template_name, data)
            elif format == DocumentFormat.PDF:
                content = await self._generate_pdf(template_name, data)
            elif format == DocumentFormat.PPTX:
                content = await self._generate_pptx(template_name, data)
            else:
                raise ValueError(f"Unsupported format: {format}")
            
            return GeneratedDocument(
                content=content,
                filename=filename,
                format=format,
                content_type=self.content_types[format],
                size=len(content),
                generated_at=datetime.now()
            )
            
        except Exception as e:
            logger.error(f"Error generating document: {str(e)}")
            raise
    
    async def _generate_html(self, template_name: str, data: Dict[str, Any]) -> bytes:
        """Generate HTML document"""
        try:
            # Load and render template
            template = self.jinja_env.get_template(f"{template_name}.html")
            
            # Add default context
            context = {
                'generated_at': datetime.now(),
                'generator': 'TenderWise AI',
                **data
            }
            
            # Render template
            html_content = template.render(context)
            
            return html_content.encode('utf-8')
            
        except Exception as e:
            logger.error(f"Error generating HTML: {str(e)}")
            raise
    
    async def _generate_pdf(self, template_name: str, data: Dict[str, Any]) -> bytes:
        """Generate PDF document from HTML template"""
        try:
            # First generate HTML
            html_content = await self._generate_html(template_name, data)
            
            # Convert HTML to PDF using WeasyPrint
            try:
                import weasyprint
                
                # Create PDF from HTML
                pdf_document = weasyprint.HTML(string=html_content.decode('utf-8'))
                pdf_bytes = pdf_document.write_pdf()
                
                return pdf_bytes
                
            except (ImportError, OSError) as e:
                # Fallback using reportlab if WeasyPrint fails
                logger.warning(f"WeasyPrint failed, using ReportLab fallback: {e}")
                return await self._generate_pdf_reportlab(data)
                
        except Exception as e:
            logger.error(f"Error generating PDF: {str(e)}")
            raise
    
    async def _generate_pdf_reportlab(self, data: Dict[str, Any]) -> bytes:
        """Fallback PDF generation using ReportLab"""
        try:
            from reportlab.lib.pagesizes import letter, A4
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.lib import colors
            
            # Create PDF buffer
            buffer = io.BytesIO()
            
            # Create document
            doc = SimpleDocTemplate(buffer, pagesize=A4)
            styles = getSampleStyleSheet()
            story = []
            
            # Title
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                spaceAfter=30,
                textColor=colors.darkblue,
                alignment=1  # Center alignment
            )
            
            title = data.get('title', 'Generated Document')
            story.append(Paragraph(title, title_style))
            story.append(Spacer(1, 20))
            
            # Content sections
            if 'sections' in data:
                for section in data['sections']:
                    # Section header
                    story.append(Paragraph(section.get('title', ''), styles['Heading2']))
                    story.append(Spacer(1, 10))
                    
                    # Section content
                    content = section.get('content', '')
                    if isinstance(content, list):
                        for item in content:
                            story.append(Paragraph(f"• {item}", styles['Normal']))
                    else:
                        story.append(Paragraph(content, styles['Normal']))
                    
                    story.append(Spacer(1, 15))
            
            # Build PDF
            doc.build(story)
            buffer.seek(0)
            
            return buffer.getvalue()
            
        except ImportError:
            raise ValueError("No PDF generation library available")
    
    async def _generate_pptx(self, template_name: str, data: Dict[str, Any]) -> bytes:
        """Generate PowerPoint presentation"""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.enum.text import PP_ALIGN
            from pptx.dml.color import RGBColor
            
            # Create presentation
            prs = Presentation()
            
            # Title slide
            slide_layout = prs.slide_layouts[0]  # Title slide layout
            slide = prs.slides.add_slide(slide_layout)
            
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            title.text = data.get('title', 'Generated Presentation')
            subtitle.text = data.get('subtitle', f"Generated on {datetime.now().strftime('%B %d, %Y')}")
            
            # Content slides
            if 'slides' in data:
                for slide_data in data['slides']:
                    self._add_content_slide(prs, slide_data)
            
            # Save to buffer
            buffer = io.BytesIO()
            prs.save(buffer)
            buffer.seek(0)
            
            return buffer.getvalue()
            
        except ImportError:
            raise ValueError("python-pptx library not available")
    
    def _add_content_slide(self, prs, slide_data: Dict[str, Any]):
        """Add content slide to presentation"""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN
        
        # Use title and content layout
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        
        # Set title
        title_shape = slide.shapes.title
        title_shape.text = slide_data.get('title', 'Slide Title')
        
        # Set content
        content_placeholder = slide.placeholders[1]
        text_frame = content_placeholder.text_frame
        text_frame.clear()
        
        content = slide_data.get('content', [])
        if isinstance(content, str):
            content = [content]
        
        for i, item in enumerate(content):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
            
            p.text = item
            p.level = slide_data.get('level', 0)
    
    def create_template(
        self,
        name: str,
        type: DocumentType,
        format: DocumentFormat,
        content: str,
        description: str = "",
        variables: Optional[Dict[str, str]] = None
    ) -> DocumentTemplate:
        """Create a new document template"""
        try:
            # Create template file
            template_path = self.template_dir / f"{name}.{format.value}"
            
            with open(template_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            template = DocumentTemplate(
                name=name,
                type=type,
                format=format,
                template_path=str(template_path),
                description=description,
                variables=variables or {}
            )
            
            logger.info(f"Created template: {name}")
            return template
            
        except Exception as e:
            logger.error(f"Error creating template: {str(e)}")
            raise
    
    def list_templates(self, type: Optional[DocumentType] = None) -> list[DocumentTemplate]:
        """List available templates"""
        templates = []
        
        for template_file in self.template_dir.glob("*.html"):
            # Extract metadata from template file
            name = template_file.stem
            
            # You could read metadata from template files or a separate config
            template = DocumentTemplate(
                name=name,
                type=DocumentType.REPORT,  # Default type
                format=DocumentFormat.HTML,
                template_path=str(template_file),
                description=f"Template: {name}",
                variables={}
            )
            
            if type is None or template.type == type:
                templates.append(template)
        
        return templates
    
    # Template filter functions
    def _currency_filter(self, value: Union[int, float], currency: str = "USD") -> str:
        """Format currency values"""
        if value is None:
            return ""
        return f"${value:,.2f} {currency}"
    
    def _datetime_filter(self, value: datetime, format: str = "%B %d, %Y") -> str:
        """Format datetime values"""
        if value is None:
            return ""
        return value.strftime(format)
    
    def _percentage_filter(self, value: Union[int, float], decimals: int = 1) -> str:
        """Format percentage values"""
        if value is None:
            return ""
        return f"{value:.{decimals}f}%"


# Global instance
document_generator = DocumentGenerator()